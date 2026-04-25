from __future__ import annotations

import hashlib
import json
from pathlib import Path
from typing import Iterable

import pandas as pd
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from ops_deck_rag.models import MetricRecord, SlideRecord


def deck_id_for(path: Path, month: str, business_unit: str | None) -> str:
    raw = f"{path.name}|{month}|{business_unit or ''}"
    return hashlib.sha1(raw.encode("utf-8")).hexdigest()[:12]


def safe_text(value: object) -> str:
    return str(value).strip() if value is not None else ""


def extract_shape_text(shape) -> str:
    if hasattr(shape, "has_text_frame") and shape.has_text_frame:
        return safe_text(shape.text)
    return ""


def extract_table_markdown(shape) -> tuple[str, list[dict[str, str]]]:
    rows = []
    for row in shape.table.rows:
        rows.append([safe_text(cell.text) for cell in row.cells])
    if not rows:
        return "", []
    df = pd.DataFrame(rows[1:], columns=rows[0]) if len(rows) > 1 else pd.DataFrame(rows)
    markdown = df.to_markdown(index=False)
    metric_rows = []
    headers = rows[0]
    for row in rows[1:]:
        row_dict = {headers[i] if i < len(headers) else f"col_{i}": row[i] for i in range(len(row))}
        metric_rows.append(row_dict)
    return markdown, metric_rows


def extract_chart_markdown(shape) -> tuple[str, list[dict[str, object]]]:
    chart = shape.chart
    title = None
    if chart.has_title and chart.chart_title and chart.chart_title.text_frame:
        title = safe_text(chart.chart_title.text_frame.text)

    rows: list[dict[str, object]] = []
    try:
        for plot in chart.plots:
            categories = []
            try:
                categories = [safe_text(c.label) for c in plot.categories]
            except Exception:
                categories = []

            for series in plot.series:
                series_name = safe_text(getattr(series, "name", "Series")) or "Series"
                values = list(getattr(series, "values", []) or [])
                for idx, value in enumerate(values):
                    category = categories[idx] if idx < len(categories) else str(idx + 1)
                    rows.append({
                        "chart_title": title or "Untitled chart",
                        "series": series_name,
                        "category": category,
                        "value": value,
                    })
    except Exception:
        return f"Chart detected: {title or 'Untitled chart'} [data extraction failed]", []

    if not rows:
        return f"Chart detected: {title or 'Untitled chart'} [no data extracted]", []
    return pd.DataFrame(rows).to_markdown(index=False), rows


def iter_shapes_recursive(shapes) -> Iterable:
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from iter_shapes_recursive(shape.shapes)
        else:
            yield shape


def extract_deck(path: Path, month: str, business_unit: str | None = None) -> tuple[list[SlideRecord], list[MetricRecord]]:
    presentation = Presentation(path)
    deck_id = deck_id_for(path, month, business_unit)
    slides: list[SlideRecord] = []
    metrics: list[MetricRecord] = []

    for slide_idx, slide in enumerate(presentation.slides, start=1):
        text_parts: list[str] = []
        tables_markdown: list[str] = []
        charts_markdown: list[str] = []
        slide_title = None

        if slide.shapes.title is not None:
            slide_title = safe_text(slide.shapes.title.text)

        for shape in iter_shapes_recursive(slide.shapes):
            shape_text = extract_shape_text(shape)
            if shape_text:
                text_parts.append(shape_text)

            if getattr(shape, "has_table", False):
                markdown, rows = extract_table_markdown(shape)
                if markdown:
                    tables_markdown.append(markdown)
                for row in rows:
                    metric_name = next(iter(row.values()), "table_row") if row else "table_row"
                    metrics.append(MetricRecord(
                        deck_id=deck_id,
                        deck_name=path.name,
                        month=month,
                        business_unit=business_unit,
                        slide_number=slide_idx,
                        slide_title=slide_title,
                        table_title=slide_title,
                        metric_name=safe_text(metric_name),
                        value=json.dumps(row),
                        source_type="table",
                    ))

            if getattr(shape, "has_chart", False):
                markdown, rows = extract_chart_markdown(shape)
                if markdown:
                    charts_markdown.append(markdown)
                for row in rows:
                    metrics.append(MetricRecord(
                        deck_id=deck_id,
                        deck_name=path.name,
                        month=month,
                        business_unit=business_unit,
                        slide_number=slide_idx,
                        slide_title=slide_title,
                        chart_title=safe_text(row.get("chart_title")),
                        metric_name=safe_text(row.get("chart_title")) or safe_text(row.get("series")),
                        category=safe_text(row.get("category")),
                        series=safe_text(row.get("series")),
                        value=row.get("value"),
                        source_type="chart",
                    ))

        slides.append(SlideRecord(
            deck_id=deck_id,
            deck_name=path.name,
            month=month,
            business_unit=business_unit,
            slide_number=slide_idx,
            slide_title=slide_title,
            text="\n".join(dict.fromkeys(text_parts)),
            tables_markdown=tables_markdown,
            charts_markdown=charts_markdown,
        ))

    return slides, metrics
